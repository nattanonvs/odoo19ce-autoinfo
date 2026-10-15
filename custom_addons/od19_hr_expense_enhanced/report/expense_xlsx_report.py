# -*- coding: utf-8 -*-
import base64
import io

from odoo import models, _

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    openpyxl = None


class ExpenseXlsxReport(models.AbstractModel):
    _name = 'report.od19_hr_expense_enhanced.expense_xlsx'
    _description = 'Expense XLSX Report'

    def create_xlsx_report(self, docids, data):
        """Generate the XLSX report and return a base64-encoded string."""
        if openpyxl is None:
            raise RuntimeError('openpyxl is not installed')

        expenses = self.env['hr.expense'].browse(docids)

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = 'Expenses'

        headers = [
            'Employee / พนักงาน',
            'Department / แผนก',
            'Expense Date / วันที่',
            'Description / รายละเอียด',
            'Category / หมวดหมู่',
            'Amount / จำนวนเงิน',
            'State / สถานะ',
            'Cash State / สถานะเงินสด',
        ]

        rows = []
        state_selection = dict(self.env['hr.expense']._fields['state'].selection or [])
        cash_selection = dict(self.env['hr.expense']._fields['od19_cash_state'].selection or [])

        for expense in expenses:
            rows.append({
                'employee': expense.employee_id.name or '',
                'department': (expense.department_id.name
                               or expense.employee_id.department_id.name or ''),
                'date': expense.date or '',
                'description': expense.name or '',
                'category': (expense.product_id.name
                             or expense.category_id.name or ''),
                'amount': expense.total_amount or 0.0,
                'state': state_selection.get(expense.state, expense.state or ''),
                'cash_state': cash_selection.get(
                    expense.od19_cash_state, expense.od19_cash_state or ''
                ),
            })

        self._write_sheet(sheet, headers, rows)

        # Save to in-memory buffer
        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return base64.b64encode(buffer.read())

    def _write_sheet(self, sheet, headers, rows):
        """Write headers and rows to a worksheet with basic formatting."""
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(
            start_color='4F81BD', end_color='4F81BD', fill_type='solid'
        )
        header_alignment = Alignment(horizontal='center', vertical='center')

        for col_idx, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

        for row_idx, row in enumerate(rows, start=2):
            sheet.cell(row=row_idx, column=1, value=row.get('employee', ''))
            sheet.cell(row=row_idx, column=2, value=row.get('department', ''))
            sheet.cell(row=row_idx, column=3, value=row.get('date', ''))
            sheet.cell(row=row_idx, column=4, value=row.get('description', ''))
            sheet.cell(row=row_idx, column=5, value=row.get('category', ''))
            sheet.cell(row=row_idx, column=6, value=row.get('amount', 0.0))
            sheet.cell(row=row_idx, column=7, value=row.get('state', ''))
            sheet.cell(row=row_idx, column=8, value=row.get('cash_state', ''))

        # Auto-size columns
        for col_idx in range(1, len(headers) + 1):
            max_length = len(str(headers[col_idx - 1]))
            for row_idx in range(2, len(rows) + 2):
                cell_value = sheet.cell(row=row_idx, column=col_idx).value
                if cell_value is not None:
                    max_length = max(max_length, len(str(cell_value)))
            sheet.column_dimensions[
                openpyxl.utils.get_column_letter(col_idx)
            ].width = min(max_length + 2, 50)
